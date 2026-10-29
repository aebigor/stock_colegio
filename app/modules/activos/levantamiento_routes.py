import re
from datetime import datetime, timezone
from io import BytesIO
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from openpyxl import load_workbook
import pdfplumber
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.templates import templates
from app.modules.auth.dependencies import get_current_user, user_area, require_roles, LOGISTICA_ROLES, SUPER_ADMIN_ROLES
from app.modules.auth.models import Usuario
from app.modules.categorias.service import CategoriaService
from app.modules.personas.service import PersonaService
from app.modules.ubicaciones.service import UbicacionService
from .levantamiento_service import LevantamientoService
from .model import Activo
from .levantamiento_model import LevantamientoDetalle
from app.modules.categorias.model import Categoria

router = APIRouter(prefix="/activos/inventario-fisico", tags=["Inventario físico"], dependencies=[Depends(get_current_user)])
public_router = APIRouter(prefix="/inventario-fisico-publico", tags=["Inventario físico público"])


def _nombre_usuario(user: Usuario) -> str:
    return " ".join(x for x in [getattr(user, "nombre", ""), getattr(user, "apellido", "")] if x).strip() or user.usuario


def _normalizar(valor):
    if valor is None:
        return ""
    texto = str(valor).strip().lower()
    reemplazos = str.maketrans("áéíóúüñ", "aeiouun")
    texto = texto.translate(reemplazos)
    return re.sub(r"[^a-z0-9]+", "", texto)


def _buscar_encabezados(ws):
    aliases = {
        "nombre": {"elemento", "nombre", "articulo", "item"},
        "referencia": {"referencia", "codigo", "codigobarras", "serial"},
        "cantidad": {"cant", "cantidad", "cantidadsistema"},
        "observacion": {"observacion", "observaciones", "detalle"},
        "estado": {"estado", "estadodelactivo"},
        "categoria": {"categoria"},
        "marca": {"marca"},
        "modelo": {"modelo"},
    }
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30)):
        found = {}
        for cell in row:
            key = _normalizar(cell.value)
            if not key:
                continue
            for campo, opciones in aliases.items():
                if key in {_normalizar(x) for x in opciones}:
                    found[campo] = cell.column
        if "nombre" in found:
            return row[0].row, found
    return None, {}


def _valor_fila(ws, row, col):
    if not col:
        return ""
    return str(ws.cell(row=row, column=col).value or "").strip()


def _parsear_cantidad(valor):
    """Convierte valores como 4, "16 CAJAS", "7 JUEGOS" o "3 UND"."""
    texto = str(valor or "").strip().replace("\ufeff", "")
    match = re.search(r"(\d+(?:[.,]\d+)?)", texto)
    if not match:
        return 0
    try:
        return max(0, int(float(match.group(1).replace(",", "."))))
    except ValueError:
        return 0


def _categoria_prefijo(nombre, db):
    base = re.sub(r"[^A-Za-z0-9]", "", nombre.upper())
    base = (base[:6] or "GEN")
    prefijo = base[:10]
    existente = db.query(Categoria).filter(Categoria.prefijo.ilike(prefijo)).first()
    if not existente:
        return prefijo
    for i in range(2, 1000):
        candidato = (base[:8] + str(i))[:10]
        if not db.query(Categoria).filter(Categoria.prefijo.ilike(candidato)).first():
            return candidato
    raise ValueError(f"No fue posible generar un prefijo para la categoría '{nombre}'.")


def _obtener_o_crear_categoria(nombre, db):
    nombre = re.sub(r"\s+", " ", str(nombre or "").strip()).strip(" .:-")
    if not nombre:
        return None
    categoria = db.query(Categoria).filter(Categoria.nombre.ilike(nombre)).first()
    if categoria:
        return categoria
    categoria = Categoria(
        nombre=nombre[:120],
        prefijo=_categoria_prefijo(nombre, db),
        descripcion="Creada automáticamente durante la importación de inventario físico.",
    )
    db.add(categoria)
    db.flush()
    return categoria


def _reconstruir_categoria_vertical(page, cell):
    """Lee categorías impresas verticalmente dentro de la columna ITEM."""
    crop = page.crop(cell)
    palabras = crop.extract_words(use_text_flow=False, keep_blank_chars=False)
    if not palabras:
        return ""
    # El PDF de inventario usa texto rotado: al extraerlo queda fragmentado
    # de arriba hacia abajo. Invertir el orden reconstruye MOBILIARIO y
    # CAJAS EQUIPOS ROBOTICA correctamente.
    piezas = [p["text"].strip() for p in sorted(palabras, key=lambda x: x["top"]) if p["text"].strip()]
    texto = "".join(piezas)[::-1]
    texto = re.sub(r"[^A-Za-z0-9ÁÉÍÓÚÜÑáéíóúüñ ]+", " ", texto)
    # Los fragmentos de texto rotado pueden venir pegados. Aplicamos los
    # nombres conocidos del formato y, si no, dejamos el texto reconstruido.
    compact = _normalizar(texto)
    conocidos = {
        "mobiliario": "MOBILIARIO",
        "cajasequiposrobotica": "CAJAS EQUIPOS ROBOTICA",
    }
    if compact in conocidos:
        return conocidos[compact]
    # Intento adicional para separar palabras comunes.
    if "CAJASEQUIPOSROBOTICA" in compact:
        return "CAJAS EQUIPOS ROBOTICA"
    if "MOBILIARIO" in compact:
        return "MOBILIARIO"
    return texto.strip()


def _filas_pdf(contenido):
    """Extrae filas del formato PDF real usado por Logística."""
    filas = []
    with pdfplumber.open(BytesIO(contenido)) as pdf:
        for page in pdf.pages:
            tablas = page.find_tables()
            if not tablas:
                continue
            tabla = tablas[0]
            # Las celdas altas de la primera columna son las categorías
            # verticales/agrupadoras del formato.
            rangos_categoria = []
            for cell in tabla.cells:
                x0, top, x1, bottom = cell
                if x0 <= tabla.bbox[0] + 1 and (bottom - top) >= 25 and top > tabla.bbox[1] + 90:
                    nombre_cat = _reconstruir_categoria_vertical(page, cell)
                    if nombre_cat and len(_normalizar(nombre_cat)) >= 4:
                        rangos_categoria.append((top, bottom, nombre_cat))

            words = page.extract_words(use_text_flow=False, keep_blank_chars=False)
            # Agrupar palabras por línea dentro de la columna ELEMENTO.
            lineas = {}
            for w in words:
                x0, x1, top = w["x0"], w["x1"], w["top"]
                if x0 < 245 or x0 >= 398 or top < 175:
                    continue
                key = round(top / 2) * 2
                lineas.setdefault(key, []).append(w)

            for top_key, ws in sorted(lineas.items()):
                ws = sorted(ws, key=lambda w: w["x0"])
                nombre = " ".join(w["text"] for w in ws).strip()
                if not nombre or _normalizar(nombre) in {"elemento", "observaciones", "observacion"}:
                    continue

                # Buscar cantidad y referencia en la misma línea por posición.
                cantidad_txt = "".join(w["text"] for w in words if 450 <= w["x0"] < 515 and abs(w["top"] - top_key) <= 2)
                referencia = " ".join(w["text"] for w in words if 398 <= w["x0"] < 453 and abs(w["top"] - top_key) <= 2).strip()
                observacion = " ".join(w["text"] for w in words if 619 <= w["x0"] < 790 and abs(w["top"] - top_key) <= 2).strip()
                estados = {
                    "Bueno": any(w["text"].strip().upper() == "B" for w in words if 515 <= w["x0"] < 550 and abs(w["top"] - top_key) <= 2),
                    "Regular": any(w["text"].strip().upper() == "R" for w in words if 550 <= w["x0"] < 585 and abs(w["top"] - top_key) <= 2),
                    "Malo": any(w["text"].strip().upper() == "M" for w in words if 585 <= w["x0"] < 620 and abs(w["top"] - top_key) <= 2),
                }
                estado = next((k for k, v in estados.items() if v), "Bueno")
                cantidad = _parsear_cantidad(cantidad_txt)
                if cantidad <= 0:
                    continue
                y = top_key
                categoria = next((cat for top, bottom, cat in rangos_categoria if top - 1 <= y <= bottom + 1), "")
                filas.append({
                    "nombre": nombre,
                    "referencia": referencia,
                    "cantidad": cantidad,
                    "categoria": categoria,
                    "estado": estado,
                    "observacion": observacion,
                })
    return filas


def _agrupar_detalles_vista(detalles):
    """Agrupa solo para presentar; conserva cada detalle/activo en la base de datos."""
    grupos = {}
    individuales = []
    for d in detalles:
        # Serial real significa identidad individual: nunca se agrupa.
        if d.serial:
            key = ("individual", d.id)
        else:
            key = (
                _normalizar(d.nombre), d.categoria_id or 0,
                _normalizar(d.marca), _normalizar(d.modelo),
                bool(d.es_nuevo), (d.estado_fisico or "").strip().casefold(),
                (d.observacion or "").strip().casefold(),
            )
        grupos.setdefault(key, []).append(d)

    resultado = []
    for detalles_grupo in grupos.values():
        primero = detalles_grupo[0]
        referencias = sorted(
            {str(d.referencia).strip() for d in detalles_grupo if d.referencia},
            key=lambda x: (re.sub(r"\d+", lambda m: m.group().zfill(12), x), x.casefold()),
        )
        if len(referencias) > 1:
            rango = f"Desde {referencias[0]} hasta {referencias[-1]}"
        elif referencias:
            rango = referencias[0]
        else:
            rango = "—"
        resultado.append({
            "detalles": detalles_grupo,
            "ids": ",".join(str(d.id) for d in detalles_grupo),
            "nombre": primero.nombre,
            "categoria": primero.categoria,
            "cantidad_sistema": sum(int(d.cantidad_sistema or 0) for d in detalles_grupo),
            "cantidad_encontrada": sum(int(d.cantidad_encontrada or 0) for d in detalles_grupo),
            "referencia_rango": rango,
            "estado_fisico": primero.estado_fisico,
            "observacion": primero.observacion,
            "es_nuevo": primero.es_nuevo,
            "cantidad": len(detalles_grupo),
            "todos_verificados": all(d.verificado_logistica for d in detalles_grupo),
            "alguno_verificado": any(d.verificado_logistica for d in detalles_grupo),
            "activo": primero.activo,
        })
    return sorted(resultado, key=lambda g: (str(g["nombre"] or "").casefold(), g["referencia_rango"]))


def _es_logistica(user: Usuario) -> bool:
    return bool(user.rol and user.rol.nombre in (LOGISTICA_ROLES | SUPER_ADMIN_ROLES))


@public_router.get("/{token}", response_class=HTMLResponse)
def docente(token: str, request: Request, db: Session = Depends(get_db)):
    service = LevantamientoService(db)
    lev = service.obtener_por_token(token)
    if not lev:
        raise HTTPException(status_code=404, detail="Enlace de inventario no encontrado.")
    # Fijamos el tenant antes de consultar categorías para mantener aislamiento.
    db.info["empresa_id"] = lev.empresa_id
    categorias = CategoriaService(db).listar(estado="Activas")
    return templates.TemplateResponse(request=request, name="activos/inventario_fisico_publico.html", context={
        "title": "Verificación de inventario", "levantamiento": lev, "categorias": categorias, "token": token
    })


@public_router.post("/{token}/detalle/{detalle_id}")
def docente_actualizar(token: str, detalle_id: int, cantidad_encontrada: int = Form(...), estado_fisico: str = Form(...), observacion: str = Form(""), db: Session = Depends(get_db)):
    lev = LevantamientoService(db).obtener_por_token(token)
    if not lev:
        raise HTTPException(status_code=404, detail="Enlace de inventario no encontrado.")
    db.info["empresa_id"] = lev.empresa_id
    try:
        LevantamientoService(db).actualizar_detalle_publico(token, detalle_id, cantidad_encontrada, estado_fisico, observacion)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/inventario-fisico-publico/{token}", status_code=303)


@public_router.post("/{token}/hallazgo")
def docente_hallazgo(token: str, nombre: str = Form(...), categoria_id: int = Form(...), cantidad: int = Form(1), estado_fisico: str = Form("Bueno"), referencia: str = Form(""), marca: str = Form(""), modelo: str = Form(""), serial: str = Form(""), observacion: str = Form(""), db: Session = Depends(get_db)):
    lev = LevantamientoService(db).obtener_por_token(token)
    if not lev:
        raise HTTPException(status_code=404, detail="Enlace de inventario no encontrado.")
    db.info["empresa_id"] = lev.empresa_id
    try:
        LevantamientoService(db).agregar_hallazgo_publico(token, nombre, categoria_id, cantidad, estado_fisico, observacion, referencia, marca, modelo, serial)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/inventario-fisico-publico/{token}", status_code=303)


@public_router.post("/{token}/enviar")
def docente_enviar(token: str, firma: str = Form(""), db: Session = Depends(get_db)):
    lev = LevantamientoService(db).obtener_por_token(token)
    if not lev:
        raise HTTPException(status_code=404, detail="Enlace de inventario no encontrado.")
    db.info["empresa_id"] = lev.empresa_id
    try:
        LevantamientoService(db).enviar(lev.id, firma=firma)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/inventario-fisico-publico/{token}", status_code=303)


@router.get("/", response_class=HTMLResponse)
def index(request: Request, levantamiento_id: int | None = None, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    # El control físico debe permitir revisar cualquier ubicación de la empresa,
    # independientemente del área del usuario que está diligenciando.
    area = user_area(user)
    es_super_admin = bool(user.rol and user.rol.nombre in SUPER_ADMIN_ROLES)
    es_logistica = _es_logistica(user)
    service = LevantamientoService(db)
    # Solo Super Admin ve ambas áreas. Logística y Medios ven su propio inventario.
    area_consulta = "Todas" if es_super_admin else area
    levantamientos = service.listar(area_consulta)
    levantamiento = service.obtener(levantamiento_id, area_consulta) if levantamiento_id else None
    if levantamiento and not es_super_admin and levantamiento.area != area:
        levantamiento = None
    ubicaciones = UbicacionService(db).listar("Todas")
    personas = PersonaService(db).listar()
    categorias = CategoriaService(db).listar(estado="Activas")
    return templates.TemplateResponse(
        request=request,
        name="activos/inventario_fisico.html",
        context={
            "title": "Inventario físico",
            "levantamientos": levantamientos,
            "levantamiento": levantamiento,
            "grupos_detalles": _agrupar_detalles_vista(levantamiento.detalles) if levantamiento else [],
            "ubicaciones": ubicaciones,
            "personas": personas,
            "categorias": categorias,
            "area": area,
            "usuario_actual": _nombre_usuario(user),
            "current_user_id": user.id,
            "es_logistica": es_logistica,
            "es_super_admin": es_super_admin,
        },
    )


@router.post("/crear")
def crear(ubicacion_id: int = Form(...), responsable_id: int | None = Form(None), db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    try:
        lev = LevantamientoService(db).crear_borrador(ubicacion_id, responsable_id, user_area(user), user.id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/?levantamiento_id={lev.id}", status_code=303)


@router.post("/detalle")
def agregar_detalle(
    levantamiento_id: int = Form(...),
    nombre: str = Form(...),
    categoria_id: int = Form(...),
    cantidad: int = Form(1),
    estado_fisico: str = Form("Bueno"),
    referencia: str = Form(""),
    marca: str = Form(""),
    modelo: str = Form(""),
    serial: str = Form(""),
    observacion: str = Form(""),
    db: Session = Depends(get_db),
    user: Usuario = Depends(get_current_user),
):
    try:
        LevantamientoService(db).agregar_hallazgo(levantamiento_id, nombre, categoria_id, cantidad, estado_fisico, observacion, referencia, marca, modelo, serial)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/?levantamiento_id={levantamiento_id}", status_code=303)


@router.post("/detalle/{detalle_id}")
def editar_detalle(detalle_id: int, cantidad_encontrada: int = Form(...), estado_fisico: str = Form(...), observacion: str = Form(""), levantamiento_id: int = Form(...), db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    try:
        LevantamientoService(db).actualizar_detalle(detalle_id, cantidad_encontrada, estado_fisico, observacion)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/?levantamiento_id={levantamiento_id}", status_code=303)


@router.post("/importar-excel")
async def importar_excel(
    ubicacion_id: int = Form(...),
    responsable_id: int | None = Form(None),
    categoria_por_defecto: int | None = Form(None),
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: Usuario = Depends(get_current_user),
):
    extension = (archivo.filename or "").lower().rsplit(".", 1)[-1]
    if extension not in {"xlsx", "xlsm", "pdf"}:
        raise HTTPException(status_code=400, detail="Sube un archivo Excel (.xlsx/.xlsm) o PDF (.pdf).")

    try:
        contenido = await archivo.read()
        if not contenido:
            raise ValueError("El archivo está vacío.")

        filas = []
        if extension in {"xlsx", "xlsm"}:
            wb = load_workbook(BytesIO(contenido), data_only=True)
            ws = wb.active
            header_row, cols = _buscar_encabezados(ws)
            if not header_row:
                raise ValueError("No encontré una fila de encabezados. Usa columnas como ELEMENTO, REFERENCIA, CANT., ESTADO y OBSERVACIÓN.")
            for row in range(header_row + 1, ws.max_row + 1):
                nombre = _valor_fila(ws, row, cols.get("nombre"))
                if not nombre:
                    continue
                filas.append({
                    "nombre": nombre,
                    "referencia": _valor_fila(ws, row, cols.get("referencia")),
                    "cantidad": _parsear_cantidad(_valor_fila(ws, row, cols.get("cantidad")) or "1"),
                    "categoria": _valor_fila(ws, row, cols.get("categoria")),
                    "estado": _valor_fila(ws, row, cols.get("estado")) or "Bueno",
                    "observacion": _valor_fila(ws, row, cols.get("observacion")),
                    "marca": _valor_fila(ws, row, cols.get("marca")),
                    "modelo": _valor_fila(ws, row, cols.get("modelo")),
                    "serial": _valor_fila(ws, row, cols.get("serial")),
                })
        else:
            filas = _filas_pdf(contenido)

        if not filas:
            raise ValueError("El archivo no produjo filas válidas. En el PDF se esperan elementos con cantidad; en Excel, una columna ELEMENTO y CANT.")

        area = user_area(user)
        service = LevantamientoService(db)
        lev = service.crear_borrador(ubicacion_id, responsable_id, area, user.id)
        creados = 0
        categoria_defecto = db.query(Categoria).filter(Categoria.id == categoria_por_defecto).first() if categoria_por_defecto else None

        # El borrador ya contiene un detalle por cada activo activo de la ubicación.
        # Indexamos esos detalles para comparar el Excel contra el inventario real,
        # no para convertir cada fila del archivo en un hallazgo nuevo.
        from collections import defaultdict
        detalles_existentes = [d for d in lev.detalles if d.activo_id and not d.es_nuevo]
        por_clave = defaultdict(list)
        por_identificador = {}
        for detalle in detalles_existentes:
            activo = detalle.activo
            clave = (
                _normalizar(detalle.nombre), detalle.categoria_id or 0,
                _normalizar(detalle.marca), _normalizar(detalle.modelo),
            )
            por_clave[clave].append(detalle)
            for identificador in (getattr(activo, "codigo", None), getattr(activo, "codigo_barras", None), getattr(activo, "serial", None)):
                if identificador:
                    por_identificador[_normalizar(identificador)] = detalle
        usados = set()

        # Consolidar renglones repetidos del archivo por características compatibles.
        grupos = {}
        for fila in filas:
            nombre = str(fila.get("nombre") or "").strip()
            if not nombre:
                continue
            referencia = str(fila.get("referencia") or "").strip()
            serial = str(fila.get("serial") or "").strip()
            marca = str(fila.get("marca") or "").strip()
            modelo = str(fila.get("modelo") or "").strip()
            obs = str(fila.get("observacion") or "").strip()
            estado = str(fila.get("estado") or "Bueno").strip() or "Bueno"
            cantidad = int(fila.get("cantidad") or 0)
            if cantidad < 1:
                continue

            categoria_texto = str(fila.get("categoria") or "").strip()
            if categoria_texto:
                categoria = _obtener_o_crear_categoria(categoria_texto, db)
                categoria_id = categoria.id if categoria else None
            elif categoria_defecto:
                categoria_id = categoria_defecto.id
            else:
                # Si el archivo no especifica categoría, primero comparar por nombre,
                # marca y modelo sin forzar una categoría GENERAL que impediría coincidir.
                categoria_id = None

            # Un serial real o código de activo identifica una unidad; una referencia
            # descriptiva/modelo común se compara por características del artículo.
            identificador = serial or referencia
            detalle_id = por_identificador.get(_normalizar(identificador)) if identificador else None
            detalle_identificado = next((d for d in detalles_existentes if d.id == detalle_id), None)
            if detalle_identificado:
                if detalle_identificado.id not in usados:
                    detalle_identificado.cantidad_encontrada = 1
                    detalle_identificado.estado_fisico = estado
                    if obs:
                        detalle_identificado.observacion = obs
                    usados.add(detalle_identificado.id)
                continue

            clave = (_normalizar(nombre), categoria_id or 0, _normalizar(marca), _normalizar(modelo))
            grupo_key = (clave, _normalizar(serial), _normalizar(referencia), _normalizar(estado), _normalizar(obs))
            if grupo_key not in grupos:
                grupos[grupo_key] = {
                    "nombre": nombre, "categoria_id": categoria_id, "referencia": referencia,
                    "serial": serial, "marca": marca, "modelo": modelo, "estado": estado,
                    "observacion": obs, "cantidad": 0, "clave": clave,
                }
            grupos[grupo_key]["cantidad"] += cantidad

        for grupo in grupos.values():
            cantidad_archivo = grupo["cantidad"]
            if grupo["clave"][1] == 0:
                candidatos = [d for d in detalles_existentes
                              if d.id not in usados
                              and _normalizar(d.nombre) == grupo["clave"][0]
                              and _normalizar(d.marca) == grupo["clave"][2]
                              and _normalizar(d.modelo) == grupo["clave"][3]]
            else:
                candidatos = [d for d in por_clave.get(grupo["clave"], []) if d.id not in usados]
            # Las unidades existentes se cuentan primero. Solo el excedente del Excel
            # se registra como hallazgo nuevo; nunca se duplica la cantidad ya registrada.
            existentes_a_contar = min(cantidad_archivo, len(candidatos))
            for detalle in candidatos[:existentes_a_contar]:
                detalle.cantidad_encontrada = 1
                detalle.estado_fisico = grupo["estado"]
                if grupo["observacion"]:
                    detalle.observacion = grupo["observacion"]
                usados.add(detalle.id)
            diferencia = max(0, cantidad_archivo - existentes_a_contar)
            if diferencia:
                categoria_nueva_id = grupo["categoria_id"]
                if not categoria_nueva_id:
                    categoria_general = _obtener_o_crear_categoria("GENERAL", db)
                    categoria_nueva_id = categoria_general.id if categoria_general else None
                service.agregar_hallazgo(
                    lev.id, grupo["nombre"], categoria_nueva_id, diferencia,
                    grupo["estado"], grupo["observacion"], grupo["referencia"],
                    grupo["marca"], grupo["modelo"], grupo["serial"],
                )
                creados += 1

        # Los activos de la ubicación que no aparecieron en el archivo quedan con
        # cantidad encontrada 0 para que Logística pueda revisar el faltante.
        for detalle in detalles_existentes:
            if detalle.id not in usados:
                detalle.cantidad_encontrada = 0
        db.commit()
        # Una importación puede coincidir completamente con activos existentes y, por
        # tanto, no generar hallazgos nuevos. Eso NO significa que el archivo sea inválido:
        # el borrador sirve para verificar existentes y registrar faltantes.
        if not usados and not creados:
            # Solo eliminar el borrador recién creado si realmente no se pudo asociar
            # ninguna fila válida a activos existentes ni crear hallazgos.
            db.delete(lev)
            db.commit()
            raise ValueError("No se pudo asociar ninguna fila del archivo con activos existentes ni crear hallazgos. Revisa nombres, referencias y cantidades.")
    except Exception as exc:
        db.rollback()
        if isinstance(exc, ValueError):
            raise HTTPException(status_code=400, detail=str(exc))
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo: {exc}")
    return RedirectResponse(url=f"/activos/inventario-fisico/?levantamiento_id={lev.id}", status_code=303)


@router.get("/{levantamiento_id}/revisar", response_class=HTMLResponse)
def revisar(levantamiento_id: int, request: Request, db: Session = Depends(get_db), user: Usuario = Depends(require_roles("Super Admin", "Superadministrador", "Administrador de Logística", "Admin de Logística"))):
    area_consulta = user_area(user)
    lev = LevantamientoService(db).obtener(levantamiento_id, area_consulta)
    if not lev:
        raise HTTPException(status_code=404, detail="Levantamiento no encontrado en tu área.")
    return templates.TemplateResponse(request=request, name="activos/inventario_fisico.html", context={
        "title": "Revisar inventario físico", "levantamientos": LevantamientoService(db).listar("Todas"),
        "levantamiento": lev, "grupos_detalles": _agrupar_detalles_vista(lev.detalles), "ubicaciones": UbicacionService(db).listar("Todas"),
        "personas": PersonaService(db).listar(), "categorias": CategoriaService(db).listar(estado="Activas"),
        "area": area_consulta, "usuario_actual": _nombre_usuario(user), "es_logistica": True, "modo_revision": True,
        "es_super_admin": bool(user.rol and user.rol.nombre in SUPER_ADMIN_ROLES), "current_user_id": user.id,
    })


@router.post("/detalle/{detalle_id}/verificar")
def verificar_detalle(detalle_id: int, levantamiento_id: int = Form(...), db: Session = Depends(get_db), user: Usuario = Depends(require_roles("Super Admin", "Superadministrador", "Administrador de Logística", "Admin de Logística"))):
    try:
        LevantamientoService(db).verificar_detalle(detalle_id, _nombre_usuario(user))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/{levantamiento_id}/revisar", status_code=303)


@router.post("/detalle/{detalle_id}/desverificar")
def desverificar_detalle(detalle_id: int, levantamiento_id: int = Form(...), db: Session = Depends(get_db), user: Usuario = Depends(require_roles("Super Admin", "Superadministrador", "Administrador de Logística", "Admin de Logística"))):
    try:
        LevantamientoService(db).desverificar_detalle(detalle_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/{levantamiento_id}/revisar", status_code=303)


@router.post("/{levantamiento_id}/grupo/verificar")
def verificar_grupo(levantamiento_id: int, detalle_ids_csv: str = Form(...), db: Session = Depends(get_db), user: Usuario = Depends(require_roles("Super Admin", "Superadministrador", "Administrador de Logística", "Admin de Logística"))):
    lev = LevantamientoService(db).obtener(levantamiento_id, user_area(user))
    if not lev or lev.estado != "Pendiente de aprobación":
        raise HTTPException(status_code=404, detail="Levantamiento pendiente no encontrado en tu área.")
    try:
        ids = {int(x) for x in detalle_ids_csv.split(",") if x.strip()}
    except ValueError:
        raise HTTPException(status_code=400, detail="La selección de elementos no es válida.")
    if not ids:
        raise HTTPException(status_code=400, detail="No se seleccionaron elementos.")
    detalles = db.query(LevantamientoDetalle).filter(
        LevantamientoDetalle.levantamiento_id == lev.id,
        LevantamientoDetalle.empresa_id == lev.empresa_id,
        LevantamientoDetalle.id.in_(ids),
    ).all()
    if len(detalles) != len(ids):
        raise HTTPException(status_code=400, detail="Uno o más elementos no pertenecen a este levantamiento.")
    marcar = not all(d.verificado_logistica for d in detalles)
    for d in detalles:
        d.verificado_logistica = marcar
        d.revisado_por = _nombre_usuario(user) if marcar else None
        d.revisado_at = datetime.now(timezone.utc) if marcar else None
    db.commit()
    return RedirectResponse(url=f"/activos/inventario-fisico/{levantamiento_id}/revisar", status_code=303)


@router.post("/{levantamiento_id}/eliminar")
def eliminar(levantamiento_id: int, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    """Elimina pruebas/borradores sin tocar activos del inventario maestro."""
    service = LevantamientoService(db)
    lev = service.obtener(levantamiento_id, "Todas")
    if not lev:
        raise HTTPException(status_code=404, detail="Levantamiento no encontrado.")
    if not _es_logistica(user) and lev.creado_por != user.id:
        raise HTTPException(status_code=403, detail="No puedes eliminar el levantamiento de otro usuario.")
    try:
        service.eliminar(levantamiento_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url="/activos/inventario-fisico/", status_code=303)


@router.post("/{levantamiento_id}/enviar")
def enviar(levantamiento_id: int, firma: str = Form(""), db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    try:
        lev = LevantamientoService(db).obtener(levantamiento_id, user_area(user))
        if not lev:
            raise ValueError("Levantamiento no encontrado.")
        if not _es_logistica(user) and lev.creado_por != user.id:
            raise ValueError("No puedes enviar el levantamiento de otro usuario.")
        LevantamientoService(db).enviar(levantamiento_id, firma=firma)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/?levantamiento_id={levantamiento_id}", status_code=303)


@router.post("/{levantamiento_id}/aprobar")
def aprobar(levantamiento_id: int, db: Session = Depends(get_db), user: Usuario = Depends(require_roles("Super Admin", "Superadministrador", "Administrador de Logística", "Admin de Logística"))):
    area = user_area(user)
    if area == "Todas":
        area = "Todas"
    lev = LevantamientoService(db).obtener(levantamiento_id, area)
    if not lev:
        raise HTTPException(status_code=404, detail="Levantamiento no encontrado.")
    try:
        creados = LevantamientoService(db).aprobar(levantamiento_id, user.id, _nombre_usuario(user))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/?levantamiento_id={levantamiento_id}", status_code=303)


@router.post("/{levantamiento_id}/devolver-todo")
def devolver_todo(levantamiento_id: int, db: Session = Depends(get_db), user: Usuario = Depends(require_roles("Super Admin", "Superadministrador", "Administrador de Logística", "Admin de Logística"))):
    area = user_area(user)
    lev = LevantamientoService(db).obtener(levantamiento_id, area)
    if not lev:
        raise HTTPException(status_code=404, detail="Levantamiento no encontrado en tu área.")
    try:
        devueltos = LevantamientoService(db).devolver_todo(levantamiento_id, _nombre_usuario(user))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return RedirectResponse(url=f"/activos/inventario-fisico/?levantamiento_id={levantamiento_id}", status_code=303)
